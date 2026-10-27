from dataclasses import dataclass
from pathlib import Path
import hashlib
import numpy as np
import openpyxl


@dataclass
class Record:
    dataset: str
    source: str
    sheet: str
    current: np.ndarray
    voltage: np.ndarray | None
    time: np.ndarray | None
    settings: dict
    group: str
    device: str
    sha256: str
    rows: int
    columns: list

    @property
    def key(self):
        return self.source + '::' + self.sheet


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_settings(sheet, names):
    result = {name: {} for name in names}
    active = None
    for row in sheet.iter_rows(values_only=True):
        values = [str(v).strip() for v in row if v is not None]
        if len(values) == 1 and values[0] in result:
            active = values[0]
        elif active and 1 < len(values) <= 8:
            result[active][values[0]] = values[-1]
    return result


def load_records(root):
    root = Path(root).resolve()
    result = []
    for path in sorted((root / 'data/raw').rglob('*.xlsx')):
        if path.name.startswith('~$'):
            continue
        rel = path.relative_to(root).as_posix()
        dataset = path.relative_to(root / 'data/raw').parts[0]
        group = path.parent.name if dataset != 'original' else 'original'
        device = ('Light_other_device' if group == 'Light' else 'R4C4') if dataset == 'r4c4' else ('R4C5' if dataset == 'r4c5' else 'unspecified')
        digest = sha256(path)
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            settings = read_settings(book['Settings'], book.sheetnames) if 'Settings' in book else {}
            for sheet in book:
                it = sheet.iter_rows(values_only=True)
                header = next(it, ())
                if 'AI' not in header or not ('Time' in header or 'AV' in header):
                    continue
                fields = [name for name in ['AI', 'AV', 'Time'] if name in header]
                indices = [header.index(name) for name in fields]
                values = []
                last_row = 1
                for rowno, row in enumerate(it, 2):
                    if not any(v is not None for v in row):
                        continue
                    try:
                        sample = [float(row[i]) for i in indices]
                    except (ValueError, TypeError, IndexError) as error:
                        raise ValueError(f'{rel}/{sheet.title}:{rowno}: nonnumeric measurement') from error
                    if not np.isfinite(sample).all():
                        raise ValueError(f'{rel}/{sheet.title}:{rowno}: nonfinite measurement')
                    values.append(sample)
                    last_row = rowno
                a = np.asarray(values, dtype=float)
                if len(a) < 3:
                    raise ValueError(f'{rel}/{sheet.title}: fewer than three samples')
                data = {name: a[:, i] for i, name in enumerate(fields)}
                t = data.get('Time')
                if t is not None and np.any(np.diff(t) <= 0):
                    raise ValueError(f'{rel}/{sheet.title}: nonincreasing time')
                result.append(Record(dataset, rel, sheet.title, data['AI'], data.get('AV'),
                                     t, settings.get(sheet.title, {}), group, device, digest, last_row, list(header)))
        finally:
            book.close()
    return sorted(result, key=lambda r: (r.dataset, r.settings.get('Last Executed', ''), r.source, r.sheet))
