"""Read Clarius workbooks without changing their contents or time order."""
from dataclasses import dataclass
from pathlib import Path
import hashlib

import numpy as np
import openpyxl


@dataclass
class Record:
    source: str
    sheet: str
    sha256: str
    current: np.ndarray
    voltage: np.ndarray
    time: np.ndarray | None
    metadata: dict

    @property
    def key(self):
        return f"{self.source}::{self.sheet}"


def _settings(sheet, names):
    result, current = {name: {} for name in names}, None
    for row in sheet.iter_rows(values_only=True):
        values = [str(v).strip() for v in row if v is not None]
        if not values or len(values) > 8:  # some exports append a flattened duplicate
            continue
        if len(values) == 1 and values[0] in result:
            current = values[0]
        elif current and len(values) > 1:
            result[current][values[0]] = values[-1]
    return result


def load_records(root):
    root = Path(root).expanduser().resolve()
    if not root.is_dir():
        raise ValueError(f"Measurement directory does not exist: {root}")
    records = []
    for path in sorted(root.rglob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        sha = hashlib.sha256(path.read_bytes()).hexdigest()
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            metadata = _settings(book["Settings"], book.sheetnames) if "Settings" in book else {}
            for sheet in book:
                rows = sheet.iter_rows(values_only=True)
                header = next(rows, ())
                if not {"AI", "AV"}.issubset(header):
                    continue
                names = ["AI", "AV"] + (["Time"] if "Time" in header else [])
                indices = [header.index(name) for name in names]
                data = []
                for rowno, row in enumerate(rows, 2):
                    if not any(v is not None for v in row):
                        continue
                    try:
                        values = [float(row[k]) for k in indices]
                    except (ValueError, TypeError, IndexError) as exc:
                        raise ValueError(f"Invalid numeric row: {path.name}/{sheet.title}:{rowno}") from exc
                    if not np.isfinite(values).all():
                        raise ValueError(f"Non-finite sample: {path.name}/{sheet.title}:{rowno}")
                    data.append(values)
                if len(data) < 3:
                    raise ValueError(f"Too few samples: {path.name}/{sheet.title}")
                a = np.asarray(data)
                t = a[:, 2] if "Time" in header else None
                if t is not None and np.any(np.diff(t) <= 0):
                    raise ValueError(f"Non-increasing time: {path.name}/{sheet.title}; split acquisition first")
                records.append(Record(str(path.relative_to(root)), sheet.title, sha,
                                      a[:, 0], a[:, 1], t, metadata.get(sheet.title, {})))
        finally:
            book.close()
    if not records:
        raise ValueError("No worksheets with AI and AV columns were found")
    return records


def causal_residual(time, current, tau_s=2.0):
    """Subtract previous EMA baseline; no future data or centered smoothing."""
    if tau_s <= 0 or not np.isfinite(tau_s):
        raise ValueError("tau_s must be finite and positive")
    t, x = np.asarray(time, float), np.asarray(current, float)
    if t.shape != x.shape or len(x) < 3 or np.any(np.diff(t) <= 0):
        raise ValueError("Need matching, ordered time/current arrays")
    baseline, residual = float(x[0]), np.zeros(len(x))
    for k in range(1, len(x)):
        residual[k] = x[k] - baseline
        baseline += -np.expm1(-(t[k] - t[k-1]) / tau_s) * residual[k]
    return residual


def acf(x, max_lag=100):
    x = np.asarray(x, float) - np.mean(x)
    denominator = np.dot(x, x)
    if denominator <= np.finfo(float).tiny:
        return np.r_[1., np.zeros(min(max_lag, len(x)-1))]
    return np.array([np.dot(x[:len(x)-k], x[k:]) / denominator
                     for k in range(min(max_lag, len(x)-1)+1)])


def characterize(record, tau_s=2.0):
    x, v = record.current, record.voltage
    compliance = record.metadata.get("Compliance")
    limit = float(compliance) if compliance is not None else None
    row = dict(source=record.source, sheet=record.sheet, sha256=record.sha256,
               n=len(x), kind="IT" if record.time is not None else "IV",
               voltage_median_V=float(np.median(v)), voltage_min_V=float(v.min()),
               voltage_max_V=float(v.max()), current_mean_mA=float(x.mean()*1e3),
               current_min_mA=float(x.min()*1e3), current_max_mA=float(x.max()*1e3),
               current_std_uA=float(x.std()*1e6), compliance_A=limit,
               near_compliance_fraction=float(np.mean(np.abs(x) >= .99*limit)) if limit else None,
               acquired=record.metadata.get("Last Executed", "unknown"),
               voltage_kind=record.metadata.get("Measure Voltage", "unknown"),
               flagged_name="dead" in record.sheet.lower())
    if record.time is not None:
        t, dt = record.time, np.diff(record.time)
        residual = causal_residual(t, x, tau_s)
        warmup = int(np.searchsorted(t-t[0], 5*tau_s))
        if warmup >= len(x)-3:
            raise ValueError(f"Trace shorter than filter warmup: {record.key}")
        r = residual[warmup:]
        rho = acf(r, min(200, len(r)//4))
        stop = next((k for k in range(1, len(rho)) if rho[k] <= 0), len(rho))
        quarter = len(x)//4
        row.update(dt_median_s=float(np.median(dt)), sample_rate_Hz=float(1/np.median(dt)),
                   dt_relative_jitter=float(np.std(dt)/np.mean(dt)),
                   duration_s=float(t[-1]-t[0]), raw_lag1=float(acf(x, 1)[1]),
                   residual_std_uA=float(r.std()*1e6), residual_lag1=float(rho[1]),
                   tau_effective_s=float(np.median(dt)*(.5+rho[1:stop].sum())),
                   acf_truncated=stop == len(rho), filter_tau_s=tau_s,
                   end_minus_start_mA=float((np.median(x[-quarter:])-np.median(x[:quarter]))*1e3),
                   mean_abs_power_mW=float(np.mean(np.abs(v*x))*1e3),
                   programmed_V_times_I_energy_J=float(np.trapezoid(np.abs(v*x), t)))
    return row


def prepare_replay(record, tau_s=2., train_fraction=.4):
    if record.time is None or np.ptp(record.voltage) > 1e-5:
        raise ValueError("Replay requires a fixed-bias time trace")
    if not 0.1 <= train_fraction <= 0.8:
        raise ValueError("train_fraction must be between 0.1 and 0.8")
    t = record.time
    split = int(len(t)*train_fraction)
    warmup = int(np.searchsorted(t-t[0], 5*tau_s))
    if split-warmup < 32:
        raise ValueError("Insufficient calibration samples after warmup")
    residual = causal_residual(t, record.current, tau_s)
    center, scale = float(np.mean(residual[warmup:split])), float(np.std(residual[warmup:split]))
    if scale <= np.finfo(float).eps:
        raise ValueError("Calibration trace has no resolved variation")
    z = (residual[split:]-center)/scale
    return dict(key=record.key, source=record.source, sheet=record.sheet, sha256=record.sha256,
                split_index=split, train_fraction=train_fraction, warmup_samples=warmup,
                filter_tau_s=tau_s, calibration_center_A=center, calibration_scale_A=scale,
                test_z=z, test_time=t[split:], test_power_W=np.abs(record.current[split:]*record.voltage[split:]),
                dt_s=float(np.median(np.diff(t))))
